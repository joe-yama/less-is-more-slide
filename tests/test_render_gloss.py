"""添え訳の描画（表紙・一言・見出し）。delta spec「添え訳の見た目」「添え訳の配置」。"""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from slidekit.manuscript import parse
from slidekit.render import render

FIXTURES = Path(__file__).parent / "fixtures"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
TOL = 2  # EMU


def slides_of(tmp_path: Path, text: str):
    out = tmp_path / "out.pptx"
    render(parse(text, FIXTURES), out)
    return list(Presentation(str(out)).slides)


def box(slide, name: str):
    found = [s for s in slide.shapes if s.name == name]
    assert len(found) == 1, [s.name for s in slide.shapes]
    return found[0]


def near(actual: int, inches: float) -> bool:
    return abs(actual - int(Inches(inches))) <= TOL


def bottom(shape) -> int:
    return shape.top + shape.height


def check_gloss_paragraph(p, text: str, pt: int) -> None:
    assert p.text == text
    assert p.alignment is not None and p.alignment.name == "LEFT"
    assert p.space_before == 0
    assert len(p.runs) == 1
    f = p.runs[0].font
    assert f.size.pt == pt
    assert f.bold is False
    assert str(f.color.rgb) == "767676"
    assert f.name == "メイリオ"
    ea = p.runs[0]._r.rPr.find(f"{{{NS_A}}}ea")
    assert ea is not None and ea.get("typeface") == "メイリオ"


def test_heading_gloss(tmp_path):
    s = slides_of(tmp_path, "# T\n---\n## Results\n^ 結果\n- a\n^ あ\n")[1]
    head = s.shapes.title
    assert near(head.height, 0.625)
    gloss = box(s, "見出しの添え訳")
    assert near(gloss.top, 1.225)
    assert near(gloss.height, 0.375)
    assert near(gloss.left, head.left / 914400)
    assert bottom(gloss) < int(Inches(1.8))
    check_gloss_paragraph(gloss.text_frame.paragraphs[0], "結果", 18)


def test_heading_without_gloss_keeps_geometry(tmp_path):
    s = slides_of(tmp_path, "# T\n---\n## Results\n- a\n")[1]
    assert near(s.shapes.title.height, 0.9)
    assert not [x for x in s.shapes if x.name == "見出しの添え訳"]


@pytest.mark.parametrize("gloss_lines", [1, 2])
def test_cover_title_gloss(tmp_path, gloss_lines):
    # 24pt の枠は 11.53 インチ = 約 34 字。2 行にするには全角 35 字以上
    gloss = "あ" * (10 if gloss_lines == 1 else 40)
    s = slides_of(tmp_path, f"# Title\n^ {gloss}\nLine\n")[0]
    g = gloss_lines * 24 * 1.5 / 72
    title = s.shapes.title
    gbox = box(s, "表紙のタイトルの添え訳")
    assert near(bottom(gbox), 3.8)
    assert near(gbox.top, 3.8 - g)
    assert near(gbox.height, g)
    assert abs(bottom(title) - gbox.top) <= TOL
    assert near(title.top, 1.6 - g)
    check_gloss_paragraph(gbox.text_frame.paragraphs[0], gloss, 24)


def test_cover_line_gloss(tmp_path):
    s = slides_of(tmp_path, "# Title\nFirst\n^ 一\nSecond\n")[0]
    sentence = box(s, "表紙の文")
    assert near(sentence.height, 2.4)
    ps = sentence.text_frame.paragraphs
    assert [p.text for p in ps] == ["First", "一", "Second"]
    check_gloss_paragraph(ps[1], "一", 12)
    assert ps[0].runs[0].font.size.pt == 20


def test_cover_without_gloss_keeps_geometry(tmp_path):
    s = slides_of(tmp_path, "# Title\nFirst\n")[0]
    assert near(box(s, "表紙の文").height, 1.8)
    assert near(s.shapes.title.top, 1.6)
    assert near(s.shapes.title.height, 2.2)
    assert not [x for x in s.shapes if x.name == "表紙のタイトルの添え訳"]


def test_statement_gloss(tmp_path):
    s = slides_of(tmp_path, "# T\n---\nHello\n^ こんにちは\n")[1]
    ps = box(s, "一言").text_frame.paragraphs
    assert [p.text for p in ps] == ["Hello", "こんにちは"]
    check_gloss_paragraph(ps[1], "こんにちは", 24)
    assert ps[0].runs[0].font.size.pt == 40


def test_statement_without_gloss_has_one_paragraph(tmp_path):
    s = slides_of(tmp_path, "# T\n---\nHello\n")[1]
    assert len(box(s, "一言").text_frame.paragraphs) == 1
