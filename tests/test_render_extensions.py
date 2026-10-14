"""改行・一言の補足・中扉・16pt の表・列の強調の描画。delta spec と design D3〜D7。"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from slidekit.manuscript import parse
from slidekit.render import render

FIXTURES = Path(__file__).parent / "fixtures"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
TOL = 2  # EMU


def slides_of(tmp_path: Path, text: str):
    out = tmp_path / "out.pptx"
    render(parse(text, FIXTURES), out)
    return list(Presentation(str(out)).slides)


def box(slide, name: str):
    found = [s for s in slide.shapes if s.name == name]
    assert len(found) == 1, [s.name for s in slide.shapes]
    return found[0]


def near(actual: int, inches: float) -> bool:
    return abs(actual - int(Inches(inches))) <= TOL


def bottom(shape) -> int:
    return shape.top + shape.height


def breaks(p) -> int:
    return len(p._p.findall(f"{{{NS_A}}}br"))


# --- 改行（design D3） ---


def test_statement_break_becomes_a_line_break_in_one_paragraph(tmp_path):
    s = slides_of(tmp_path, "# T\n---\n来期は値上げせず、<br>品目を減らす。\n")[1]
    (p,) = box(s, "一言").text_frame.paragraphs
    assert breaks(p) == 1
    assert [r.text for r in p.runs] == ["来期は値上げせず、", "品目を減らす。"]


def test_table_cell_break_becomes_a_line_break(tmp_path):
    text = "# T\n---\n## 表\n| 案 | 内容 |\n|---|---|\n| C | 集計基盤で実行<br>（提案） |\n"
    s = slides_of(tmp_path, text)[1]
    table = box(s, "表").table
    (p,) = table.cell(1, 1).text_frame.paragraphs
    assert breaks(p) == 1
    assert [r.text for r in p.runs] == ["集計基盤で実行", "（提案）"]
    assert breaks(table.cell(1, 0).text_frame.paragraphs[0]) == 0


# --- 一言の補足（design D4） ---


def test_statement_notes_follow_statement_and_gloss(tmp_path):
    text = (
        "# T\n---\n結論は C 案。\n^ We go with C.\n"
        "費用は 3 案で最も小さい。\n^ Lowest cost.\n条件は人員の確保。\n"
    )
    s = slides_of(tmp_path, text)[1]
    ps = box(s, "一言").text_frame.paragraphs
    assert [p.text for p in ps] == [
        "結論は C 案。",
        "We go with C.",
        "費用は 3 案で最も小さい。",
        "Lowest cost.",
        "条件は人員の確保。",
    ]
    assert ps[0].runs[0].font.size.pt == 40
    for i in (2, 4):
        f = ps[i].runs[0].font
        assert f.size.pt == 24
        assert str(f.color.rgb) == "6B6B6B"
        assert ps[i].space_before == Pt(12)
    g = ps[3]
    assert g.runs[0].font.size.pt == 14
    assert str(g.runs[0].font.color.rgb) == "767676"
    assert g.space_before == 0


def test_statement_without_notes_is_unchanged(tmp_path):
    s = slides_of(tmp_path, "# T\n---\n結論は C 案。\n")[1]
    tf = box(s, "一言")
    (p,) = tf.text_frame.paragraphs
    assert p.text == "結論は C 案。"
    assert p.space_before is None
    assert near(tf.top, 0.6) and near(tf.height, 6.3)


# --- 中扉（design D5） ---


def test_section_title_matches_cover_title_and_has_page_number(tmp_path):
    slides = slides_of(tmp_path, "# 表題\n---\n## 見出し\n- a\n---\n# 補足資料\n---\n## 次\n- b\n")
    cover, sec = slides[0], slides[2]
    title, ctitle = sec.shapes.title, cover.shapes.title
    assert title.text_frame.text == "補足資料"
    assert (title.left, title.top, title.width, title.height) == (
        ctitle.left,
        ctitle.top,
        ctitle.width,
        ctitle.height,
    )
    assert title.text_frame.vertical_anchor == ctitle.text_frame.vertical_anchor
    assert box(sec, "ページ番号").text_frame.text == "3"
    assert [p.name for p in cover.shapes if p.name == "ページ番号"] == []


def test_section_title_gloss_is_24pt_and_ends_at_3_8_inches(tmp_path):
    slides = slides_of(tmp_path, "# 表題\n---\n# 補足資料\n^ Appendix\n")
    sec = slides[1]
    gloss = box(sec, "表紙のタイトルの添え訳")
    p = gloss.text_frame.paragraphs[0]
    assert p.text == "Appendix"
    assert p.runs[0].font.size.pt == 24
    assert near(bottom(gloss), 3.8)
    assert bottom(sec.shapes.title) == gloss.top
