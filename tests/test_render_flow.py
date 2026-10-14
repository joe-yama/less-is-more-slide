"""構成図の描画（箱・範囲の枠・注記）。delta spec「構成図の見た目」「構成図のまとまりと代替テキスト」など。"""

import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from slidekit.manuscript import parse
from slidekit.render import render

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}
TOL = 2  # EMU
FONT = "メイリオ"

FLOW = (
    "| 受付 | 保存先 | 集計 |\n"
    "- 保存先: key=V1/\n"
    "- 保存先: **key=V/**\n"
    "[社内] 受付, 保存先\n"
    "受付 -> 保存先: 送る\n"
)
PLAIN = "| 甲 | 乙 |\n甲 -> 乙\n"
NOTE = "保存先は直接参照している。"
FIGURE = f"## 今の構成\n![今の構成](構成.flow)\n{NOTE}\n"
TWO = "## 比べる\n![今](左.flow)\n|||\n![案](右.flow)\n"


def build(tmp_path: Path, body: str, **flows: str):
    for name, text in flows.items():
        (tmp_path / f"{name}.flow").write_text(text, encoding="utf-8")
    out = tmp_path / "out.pptx"
    render(parse("# 題\n---\n" + body, tmp_path), out)
    return out


def slide_of(path: Path):
    return list(Presentation(str(path)).slides)[1]


def group_of(slide):
    groups = [s for s in slide.shapes if s.shape_type == 6]
    assert len(groups) == 1, [s.name for s in slide.shapes]
    return groups[0]


def near(actual: int, inches: float) -> bool:
    return abs(actual - int(Inches(inches))) <= TOL


def kids(group, name: str):
    return [s for s in group.shapes if s.name == name]


@pytest.fixture(scope="module")
def drawn(tmp_path_factory):
    path = build(tmp_path_factory.mktemp("flow"), FIGURE, 構成=FLOW)
    return path, slide_of(path)


def test_group_shape_is_single_with_name_and_alt(drawn):
    _, slide = drawn
    g = group_of(slide)
    assert g.name == "構成図"
    assert g._element.xpath("./p:nvGrpSpPr/p:cNvPr")[0].get("descr") == "今の構成"
    assert near(g.left, 0.9) and near(g.top, 1.8)
    # 箱・枠・矢印・ラベルはすべてグループの中にある
    assert [s.name for s in slide.shapes if s.name in ("箱", "範囲の枠", "矢印")] == []


def test_two_column_right_group_is_at_right_column(tmp_path):
    path = build(tmp_path, TWO, 左=PLAIN, 右=PLAIN)
    slide = slide_of(path)
    groups = [s for s in slide.shapes if s.shape_type == 6]
    assert len(groups) == 2
    xs = sorted((g.left, g.top) for g in groups)
    assert near(xs[0][0], 0.9) and near(xs[1][0], 0.9 + 5.5 + 0.53)
    assert all(near(top, 1.8) for _, top in xs)
    assert {g._element.xpath("./p:nvGrpSpPr/p:cNvPr")[0].get("descr") for g in groups} == {
        "今",
        "案",
    }


def test_boxes_are_plain_rectangles_with_no_fill_and_ink_line(drawn):
    _, slide = drawn
    boxes = kids(group_of(slide), "箱")
    assert len(boxes) == 3
    for b in boxes:
        el = b._element
        assert el.xpath("./p:spPr/a:prstGeom/@prst") == ["rect"]
        assert el.xpath("./p:spPr/a:noFill")
        ln = el.xpath("./p:spPr/a:ln")[0]
        assert ln.get("w") == "15875"  # 1.25pt
        assert ln.xpath("./a:solidFill/a:srgbClr/@val") == ["1A1A1A"]


def test_box_text_sizes_colors_and_left_alignment(drawn):
    _, slide = drawn
    boxes = {b.text_frame.paragraphs[0].text: b for b in kids(group_of(slide), "箱")}
    store = boxes["保存先"]
    paras = store.text_frame.paragraphs
    assert [p.text for p in paras] == ["保存先", "key=V1/", "key=V/"]
    name_run = paras[0].runs[0]
    assert name_run.font.size.pt == 16 and str(name_run.font.color.rgb) == "1A1A1A"
    plain_run = paras[1].runs[0]
    assert plain_run.font.size.pt == 12 and str(plain_run.font.color.rgb) == "6B6B6B"
    emph = paras[2].runs[0]
    assert emph.font.size.pt == 12 and str(emph.font.color.rgb) == "B23A2E"
    for b in boxes.values():
        for p in b.text_frame.paragraphs:
            assert p._p.xpath("./a:pPr/@algn") == ["l"]
            for r in p.runs:
                assert r.font.name == FONT
        assert b.text_frame.vertical_anchor == 3  # MSO_ANCHOR.MIDDLE


def test_frame_is_dashed_light_gray_rectangle_with_label(drawn):
    _, slide = drawn
    g = group_of(slide)
    frames = kids(g, "範囲の枠")
    assert len(frames) == 1
    el = frames[0]._element
    assert el.xpath("./p:spPr/a:prstGeom/@prst") == ["rect"]
    assert el.xpath("./p:spPr/a:noFill")
    ln = el.xpath("./p:spPr/a:ln")[0]
    assert ln.get("w") == "12700"
    assert ln.xpath("./a:solidFill/a:srgbClr/@val") == ["BDBDBD"]
    assert ln.xpath("./a:prstDash/@val") == ["dash"]
    labels = kids(g, "範囲の枠のラベル")
    assert len(labels) == 1
    p = labels[0].text_frame.paragraphs[0]
    assert p.text == "社内" and p.runs[0].font.size.pt == 12
    assert str(p.runs[0].font.color.rgb) == "6B6B6B"
    assert p._p.xpath("./a:pPr/@algn") == ["l"]


def test_bdbdbd_is_used_only_by_frame_lines(drawn):
    path, slide = drawn
    with zipfile.ZipFile(path) as z:
        xml = z.read("ppt/slides/slide2.xml").decode("utf-8")
    assert len(re.findall("BDBDBD", xml)) == 1
    hits = slide._element.xpath('.//a:srgbClr[@val="BDBDBD"]')
    assert len(hits) == 1
    hit = hits[0]
    assert hit.getparent().getparent().tag.endswith("}ln")
    shape = hit.getparent().getparent().getparent().getparent()
    assert shape.xpath("./p:nvSpPr/p:cNvPr/@name") == ["範囲の枠"]


def test_no_style_or_effects_on_any_shape(drawn):
    _, slide = drawn
    el = slide._element
    assert not el.xpath("//p:style")
    assert not el.xpath("//a:effectLst | //a:effectDag | //a:gradFill | //a:scene3d | //a:sp3d")
    assert not el.xpath('//a:prstGeom[@prst="roundRect"]')


def test_all_paragraphs_in_the_diagram_are_left_aligned(drawn):
    _, slide = drawn
    g = group_of(slide)
    paras = g._element.xpath(".//a:p")
    assert paras
    for p in paras:
        assert p.xpath("./a:pPr/@algn") == ["l"]


def test_note_is_16pt_ink_and_0_3_inch_below_the_diagram(drawn):
    _, slide = drawn
    g = group_of(slide)
    notes = [s for s in slide.shapes if s.name == "図の説明"]
    assert len(notes) == 1
    note = notes[0]
    assert abs(note.top - (g.top + g.height + int(Inches(0.3)))) <= TOL
    assert near(note.left, 0.9) and near(note.width, 11.53)
    p = note.text_frame.paragraphs[0]
    assert p.text == NOTE
    assert p.runs[0].font.size.pt == 16 and str(p.runs[0].font.color.rgb) == "1A1A1A"
    assert p._p.xpath("./a:pPr/@algn") == ["l"]
