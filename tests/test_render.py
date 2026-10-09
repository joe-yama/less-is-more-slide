import math
import re
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from slidekit.manuscript import parse
from slidekit.render import render

FIXTURES = Path(__file__).parent / "fixtures"
FONT = "メイリオ"
COLORS = {"FFFFFF", "1A1A1A", "6B6B6B", "B23A2E"}
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}

ALL_LAYOUTS = """# 四半期の振り返り
営業部 山田
2026 年 10 月
---
結論から言うと、来期は**値上げ**しない。
---
## 今期やったこと
- 新規の顧客を **3 割** 増やした
- 問い合わせの返信を 1 日以内にした
---
## 構成
![歯車](icon:gear)
毎晩 2 時に動く。
---
## 地域別の売上
| 地域 | 売上 | 前年比 |
|---|---|---|
| 東日本 | 120 | 1.1 |
| 西日本 | 90 | 0.9 |
| 九州 | 40 | 1.3 |
---
## 前と後
- 手作業で 3 時間
- 週に 1 回
|||
![ロボット](robot.txt)
"""


def build(tmp_path: Path, text: str = ALL_LAYOUTS) -> Path:
    out = tmp_path / "out.pptx"
    render(parse(text, FIXTURES), out)
    return out


@pytest.fixture(scope="module")
def pptx_path(tmp_path_factory) -> Path:
    return build(tmp_path_factory.mktemp("render"))


def part(path: Path, name: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf-8")


def slide_xmls(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        names = sorted(
            (n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.findall(r"\d+", n)[0]),
        )
        return [z.read(n).decode("utf-8") for n in names]


def all_paragraph_texts(slide) -> list[str]:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out += [p.text for p in sh.text_frame.paragraphs]
    return out


# --- 配色と書体 ---


def test_colors_are_only_the_four(pptx_path):
    used = set()
    for xml in slide_xmls(pptx_path):
        assert "schemeClr" not in xml and "sysClr" not in xml and "prstClr" not in xml
        used |= set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml))
    assert used and {c.upper() for c in used} <= COLORS


def test_background_is_white_on_every_slide(pptx_path):
    for xml in slide_xmls(pptx_path):
        m = re.search(r"<p:bg>.*?</p:bg>", xml, re.DOTALL)
        assert m, "背景を明示していない"
        assert re.findall(r'srgbClr val="(\w+)"', m.group(0)) == ["FFFFFF"]


def test_every_run_has_meiryo_latin_and_ea_size_and_color(pptx_path):
    for xml in slide_xmls(pptx_path):
        rprs = re.findall(
            r"<a:(?:rPr|endParaRPr)\b[^>]*?(?:/>|>.*?</a:(?:rPr|endParaRPr)>)", xml, re.DOTALL
        )
        assert rprs
        for r in rprs:
            assert f'<a:latin typeface="{FONT}"' in r
            assert f'<a:ea typeface="{FONT}"' in r
            assert 'sz="' in r
            assert "srgbClr" in r


def test_every_paragraph_run_is_explicit(pptx_path):
    # 文字を持つ run はすべて rPr を持つ（既定の書体に落ちない）
    for xml in slide_xmls(pptx_path):
        for r in re.findall(r"<a:r>.*?</a:r>", xml, re.DOTALL):
            assert "<a:rPr" in r


def test_theme_fonts_are_meiryo(pptx_path):
    theme = part(pptx_path, "ppt/theme/theme1.xml")
    for kind in ("majorFont", "minorFont"):
        block = re.search(rf"<a:{kind}>.*?</a:{kind}>", theme, re.DOTALL).group(0)
        assert f'<a:latin typeface="{FONT}"' in block
        assert f'<a:ea typeface="{FONT}"' in block
        assert f'<a:font script="Jpan" typeface="{FONT}"' in block


# --- 装飾を持たない ---


def test_no_decoration_elements(pptx_path):
    banned = (
        "gradFill",
        "effectLst",
        "effectDag",
        "outerShdw",
        "innerShdw",
        "prstShdw",
        "glow",
        "softEdge",
        "reflection",
        "p:timing",
        "p:transition",
        "a:tile",
        "custGeom",
        "a:scene3d",
        "a:sp3d",
    )
    for xml in slide_xmls(pptx_path):
        for tag in banned:
            assert tag not in xml, tag
        bg = re.search(r"<p:bg>.*?</p:bg>", xml, re.DOTALL).group(0)
        assert "blip" not in bg and "gradFill" not in bg


def test_shapes_are_only_text_boxes_pictures_and_tables(pptx_path):
    prs = Presentation(pptx_path)
    for slide in prs.slides:
        for sh in slide.shapes:
            el = sh._element.tag.split("}")[1]
            if el == "sp":
                is_title = sh.is_placeholder
                is_box = sh._element.xpath("./p:nvSpPr/p:cNvSpPr/@txBox") == ["1"]
                assert is_title or is_box
                assert not sh._element.xpath("./p:spPr/a:prstGeom[@prst!='rect']")
            else:
                assert el in ("pic", "graphicFrame")
                if el == "graphicFrame":
                    assert sh.has_table


def test_paragraphs_left_aligned_except_page_number(pptx_path):
    for i, xml in enumerate(slide_xmls(pptx_path), start=1):
        algns = re.findall(r'<a:pPr[^>]*?algn="(\w+)"', xml)
        assert "ctr" not in algns and "just" not in algns
        if i == 1:
            assert set(algns) <= {"l"}
        else:
            assert algns.count("r") == 1  # ページ番号だけ
    prs = Presentation(pptx_path)
    for slide in list(prs.slides)[1:]:
        for sh in slide.shapes:
            if sh.has_text_frame:
                is_num = sh.name == "ページ番号"
                for p in sh.text_frame.paragraphs:
                    a = p._p.xpath("./a:pPr/@algn")
                    assert a == (["r"] if is_num else ["l"])


# --- 体裁 ---


def test_16_9_and_slide_count(pptx_path):
    prs = Presentation(pptx_path)
    assert prs.slide_width == Inches(13.333) or abs(prs.slide_width - Inches(13.333)) < 2000
    assert prs.slide_height == Inches(7.5)
    assert len(prs.slides) == 6


def test_titles_are_title_placeholders(pptx_path):
    prs = Presentation(pptx_path)
    titles = [s.shapes.title.text if s.shapes.title is not None else None for s in prs.slides]
    assert titles == [
        "四半期の振り返り",
        None,
        "今期やったこと",
        "構成",
        "地域別の売上",
        "前と後",
    ]


def test_title_position_and_unused_placeholders_removed(pptx_path):
    prs = Presentation(pptx_path)
    for slide in prs.slides:
        t = slide.shapes.title
        if t is not None:
            assert t.left == Inches(0.9)
            assert t.top == Inches(0.6) or slide is prs.slides[0]
        for ph in slide.placeholders:
            assert ph.has_text_frame and ph.text_frame.text, "空のプレースホルダが残っている"


def test_page_numbers_bottom_right_gray_small(pptx_path):
    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    assert not any(sh.name == "ページ番号" for sh in slides[0].shapes)
    for n, slide in enumerate(slides[1:], start=2):
        boxes = [sh for sh in slide.shapes if sh.name == "ページ番号"]
        assert len(boxes) == 1
        box = boxes[0]
        assert box.text_frame.text == str(n)
        assert box.left + box.width <= prs.slide_width
        assert box.left + box.width > prs.slide_width - Inches(1.3)
        assert box.top > prs.slide_height - Inches(1.0)
        run = box.text_frame.paragraphs[0].runs[0]
        assert run.font.size.pt == 12
        assert str(run.font.color.rgb) == "6B6B6B"


def test_document_properties(pptx_path):
    cp = Presentation(pptx_path).core_properties
    assert cp.title == "四半期の振り返り"
    for f in ("author", "last_modified_by", "subject", "keywords", "comments", "category"):
        assert getattr(cp, f) == ""
    with zipfile.ZipFile(pptx_path) as z:
        for name in z.namelist():
            if name.startswith("docProps/"):
                assert b"Steve Canny" not in z.read(name)


def test_no_python_pptx_text_in_docprops(pptx_path):
    with zipfile.ZipFile(pptx_path) as z:
        for name in z.namelist():
            if name.startswith("docProps/") and name.endswith(".xml"):
                assert "python-pptx" not in z.read(name).decode("utf-8").lower()


def test_created_and_modified_are_generation_time(tmp_path):
    before = datetime.now(UTC).replace(tzinfo=None, microsecond=0)
    out = build(tmp_path)
    after = datetime.now(UTC).replace(tzinfo=None)
    cp = Presentation(out).core_properties
    assert before <= cp.created <= after
    assert before <= cp.modified <= after


# --- 強調 ---


def test_emphasis_is_color_not_bold(pptx_path):
    prs = Presentation(pptx_path)
    emphasized = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    assert not r.font.bold
                    if r.font.color.type is not None and str(r.font.color.rgb) == "B23A2E":
                        emphasized.append(r.text)
    assert emphasized == ["値上げ", "3 割"]


# --- 表 ---


def table_of(path: Path):
    prs = Presentation(path)
    slide = list(prs.slides)[4]
    return next(sh for sh in slide.shapes if sh.has_table)


def test_table_has_no_style_banding_or_fill(pptx_path):
    gf = table_of(pptx_path)
    tbl = gf.table
    assert len(tbl.rows) == 4 and len(tbl.columns) == 3
    assert tbl.first_row is False and tbl.horz_banding is False
    assert tbl._tbl.xpath("./a:tblPr/a:tableStyleId/text()") == [
        "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    ]
    for tc in tbl._tbl.xpath(".//a:tc"):
        assert tc.xpath("./a:tcPr/a:noFill")
        assert not tc.xpath("./a:tcPr/a:solidFill")


def test_table_lines_only_under_header_and_last_row(pptx_path):
    tbl = table_of(pptx_path).table
    rows = tbl._tbl.xpath("./a:tr")
    for ri, tr in enumerate(rows):
        for tc in tr.xpath("./a:tc"):
            for side in ("lnL", "lnR", "lnT"):
                assert not tc.xpath(f"./a:tcPr/a:{side}/a:solidFill"), side
            solid = tc.xpath("./a:tcPr/a:lnB/a:solidFill/a:srgbClr/@val")
            width = tc.xpath("./a:tcPr/a:lnB/@w")
            if ri in (0, len(rows) - 1):
                assert solid == ["6B6B6B"] and width == ["12700"]
            else:
                assert solid == []


def test_table_header_row_is_gray(pptx_path):
    tbl = table_of(pptx_path).table
    for ci in range(3):
        r = tbl.cell(0, ci).text_frame.paragraphs[0].runs[0]
        assert str(r.font.color.rgb) == "6B6B6B"
        r = tbl.cell(1, ci).text_frame.paragraphs[0].runs[0]
        assert str(r.font.color.rgb) == "1A1A1A"


# --- 画像 ---


def test_picture_alt_text_and_size(pptx_path):
    prs = Presentation(pptx_path)
    pics = {}
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.shape_type == 13:
                pics[sh._element.xpath("./p:nvPicPr/p:cNvPr/@descr")[0]] = (slide, sh)
    assert set(pics) == {"歯車", "ロボット"}
    for _, sh in pics.values():
        assert sh.width <= Inches(4.6) + 1 and sh.height <= Inches(4.6) + 1
        w, h = sh.image.size
        assert w == h  # アイコンと見本の格子は正方形
        dots = {"歯車": 16, "ロボット": 8}[sh._element.xpath("./p:nvPicPr/p:cNvPr/@descr")[0]]
        # design.md: 長辺が 1024 以上になる整数倍、ただし 16 倍以上 64 倍以下
        assert w == dots * min(64, max(16, math.ceil(1024 / dots)))
    # 図のレイアウト: 画像は左、説明文は右
    _, gear = pics["歯車"]
    assert gear.left == Inches(0.9)


def test_figure_caption_right_of_picture(pptx_path):
    slide = list(Presentation(pptx_path).slides)[3]
    pic = next(s for s in slide.shapes if s.shape_type == 13)
    cap = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "毎晩 2 時に動く。"]
    assert cap and cap[0].left >= pic.left + pic.width


# --- レイアウト ---


def test_statement_is_vertically_centered_in_a_text_box(pptx_path):
    slide = list(Presentation(pptx_path).slides)[1]
    boxes = [s for s in slide.shapes if s.has_text_frame and "値上げ" in s.text_frame.text]
    assert len(boxes) == 1
    assert boxes[0]._element.xpath("./p:txBody/a:bodyPr/@anchor") == ["ctr"]
    assert boxes[0].text_frame.paragraphs[0].runs[0].font.size.pt == 40


def test_bullets_use_bullet_format_not_characters(pptx_path):
    slide = list(Presentation(pptx_path).slides)[2]
    box = next(s for s in slide.shapes if s.has_text_frame and "新規" in s.text_frame.text)
    assert next(p.text for p in box.text_frame.paragraphs).startswith("新規")
    for p in box.text_frame.paragraphs:
        assert p._p.xpath("./a:pPr/a:buChar/@char") == ["・"]
        assert p.runs[0].font.size.pt == 24
    assert box.top == Inches(1.8)


def test_two_columns_geometry(pptx_path):
    slide = list(Presentation(pptx_path).slides)[5]
    left = next(s for s in slide.shapes if s.has_text_frame and "手作業" in s.text_frame.text)
    pic = next(s for s in slide.shapes if s.shape_type == 13)
    assert left.left == Inches(0.9) and left.width == Inches(5.5)
    assert pic.left >= Inches(0.9 + 5.5 + 0.53) - 1
    for p in left.text_frame.paragraphs:
        assert p.runs[0].font.size.pt == 22


def test_two_columns_text_block(tmp_path):
    out = build(tmp_path, "# 題\n---\n## 見出し\n左の文。\n|||\n- 右の項目\n")
    slide = list(Presentation(out).slides)[1]
    texts = all_paragraph_texts(slide)
    assert "左の文。" in texts and "右の項目" in texts


def test_cover_without_subtitle_lines(tmp_path):
    out = build(tmp_path, "# 題だけ\n")
    prs = Presentation(out)
    assert len(prs.slides) == 1
    assert prs.core_properties.title == "題だけ"


def test_no_leftover_parts(pptx_path):
    with zipfile.ZipFile(pptx_path) as z:
        assert "docProps/thumbnail.jpeg" not in z.namelist()
