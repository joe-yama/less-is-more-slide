"""検証済みの原稿 → PPTX。ここでは検証しない（検証は manuscript.parse の仕事）。"""

import io
import math
import re
from datetime import UTC, datetime
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

from slidekit.fit import (
    BODY_BOX_H,
    BULLET_GAP_PT,
    CELL_MARGIN_V,
    PT_BULLET,
    PT_CAPTION,
    PT_CELL,
    PT_COLUMN_BULLET,
    PT_COLUMN_TEXT,
    PT_COVER_LINE,
    PT_COVER_TITLE,
    PT_HEADING,
    PT_PAGE_NUMBER,
    PT_STATEMENT,
    TABLE_ROW_H,
)
from slidekit.manuscript import (
    Block,
    Bullets,
    BulletsBlock,
    Cover,
    Deck,
    Figure,
    ImageBlock,
    Statement,
    Table,
    Text,
    TextBlock,
    TwoColumn,
)
from slidekit.palette import ACCENT, GRAY, INK, PAPER
from slidekit.png import encode_png

FONT = "メイリオ"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_TABLE_STYLE_NONE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

SLIDE_W = Emu(12192000)  # 13.333 インチ
SLIDE_H = Emu(6858000)  # 7.5 インチ
MARGIN_X = 0.9
HEAD_TOP = 0.6
BODY_TOP = 1.8
BODY_W = 11.53
INDENT = 0.4
IMAGE_MAX = 4.6
COLUMN_W = 5.5
COLUMN_GAP = 0.53
CAPTION_W = 5.2
CAPTION_GAP = 0.5
LAYOUT_TITLE_ONLY = 5
LAYOUT_BLANK = 6
MIN_SCALE, MAX_SCALE, TARGET_PX = 1, 64, 1024
IMAGE_TARGET = 2.4
DOT_MIN, DOT_MAX = 0.05, 0.075  # 1 ドットの寸法（インチ）の下限と上限


def image_long_side(longest: int) -> float:
    """長辺のドット数から、画像の長辺（インチ）を決める。"""
    size = min(max(IMAGE_TARGET, longest * DOT_MIN), longest * DOT_MAX)
    return min(size, IMAGE_MAX)


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def _style_run(run, pt: float, color: str) -> None:
    f = run.font
    f.size = Pt(pt)
    f.bold = False
    f.color.rgb = _rgb(color)
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(f"{{{_NS_A}}}latin")
    ea = etree.SubElement(rpr, f"{{{_NS_A}}}ea")
    ea.set("typeface", FONT)
    latin.addnext(ea)


def _fill_paragraph(p, text: Text, pt: float, color: str) -> None:
    p.alignment = PP_ALIGN.LEFT
    for r in text:
        run = p.add_run()
        run.text = r.text
        _style_run(run, pt, ACCENT if r.emphasis else color)


def _bulletize(p) -> None:
    ppr = p._p.get_or_add_pPr()
    marl = Inches(INDENT)
    ppr.set("marL", str(int(marl)))
    ppr.set("indent", str(-int(marl)))
    ppr.append(parse_xml(f'<a:buFont xmlns:a="{_NS_A}" typeface="{FONT}"/>'))
    ppr.append(parse_xml(f'<a:buChar xmlns:a="{_NS_A}" char="・"/>'))


def _frame(tf, anchor: MSO_ANCHOR) -> None:
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor


def _textbox(slide, name, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    _frame(box.text_frame, anchor)
    return box.text_frame


def _paragraphs(tf, items: list[Text], pt: float, color: str, bullets: bool, gap: float = 8):
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _fill_paragraph(p, text, pt, color)
        if i > 0:
            p.space_before = Pt(gap)
        if bullets:
            _bulletize(p)


def _title(slide, text: str, x, y, w, h, pt: float, anchor=MSO_ANCHOR.TOP) -> None:
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = t.text_frame
    _frame(tf, anchor)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    _style_run(run, pt, INK)


def _page_number(slide, n: int) -> None:
    w = Inches(1.0)
    box = slide.shapes.add_textbox(SLIDE_W - Inches(MARGIN_X) - w, Inches(6.85), w, Inches(0.4))
    box.name = "ページ番号"
    _frame(box.text_frame, MSO_ANCHOR.TOP)
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = str(n)
    _style_run(run, PT_PAGE_NUMBER, GRAY)


def _picture(slide, grid, alt: str, x: float, y: float) -> None:
    longest = max(grid.width, grid.height)
    scale = min(MAX_SCALE, max(MIN_SCALE, math.ceil(TARGET_PX / longest)))
    png = encode_png(grid, scale)
    unit = image_long_side(longest) / longest
    pic = slide.shapes.add_picture(
        io.BytesIO(png),
        Inches(x),
        Inches(y),
        Inches(unit * grid.width),
        Inches(unit * grid.height),
    )
    pic.name = "図"
    pic._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", alt)


def _white_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(PAPER)
    # python-pptx が足す空の効果指定は任意の要素なので取り除く（効果を一切出さない）
    for effect in slide.background._cSld.xpath("./p:bg/p:bgPr/a:effectLst"):
        effect.getparent().remove(effect)


# --- レイアウトごと ---


def _cover(slide, s: Cover) -> None:
    _title(slide, s.title, MARGIN_X, 1.6, BODY_W, 2.2, PT_COVER_TITLE, MSO_ANCHOR.BOTTOM)
    if s.lines:
        tf = _textbox(slide, "表紙の文", MARGIN_X, 4.1, BODY_W, 1.8)
        _paragraphs(tf, s.lines, PT_COVER_LINE, GRAY, False)


def _statement(slide, s: Statement) -> None:
    tf = _textbox(slide, "一言", MARGIN_X, HEAD_TOP, BODY_W, 7.5 - 2 * HEAD_TOP, MSO_ANCHOR.MIDDLE)
    _paragraphs(tf, [s.text], PT_STATEMENT, INK, False)


def _bullets(slide, s: Bullets) -> None:
    tf = _textbox(slide, "箇条書き", MARGIN_X, BODY_TOP, BODY_W, BODY_BOX_H)
    _paragraphs(tf, s.items, PT_BULLET, INK, True, gap=BULLET_GAP_PT)


def _figure(slide, s: Figure) -> None:
    _picture(slide, s.grid, s.alt, MARGIN_X, BODY_TOP)
    if s.caption is not None:
        x = MARGIN_X + IMAGE_MAX + CAPTION_GAP
        tf = _textbox(slide, "図の説明", x, BODY_TOP, CAPTION_W, 2.4)
        _paragraphs(tf, [s.caption], PT_CAPTION, INK, False)


def _line_xml(side: str, visible: bool) -> str:
    if not visible:
        return f'<a:{side} xmlns:a="{_NS_A}" w="12700"><a:noFill/></a:{side}>'
    return (
        f'<a:{side} xmlns:a="{_NS_A}" w="12700" cap="flat" cmpd="sng" algn="ctr">'
        f'<a:solidFill><a:srgbClr val="{GRAY}"/></a:solidFill>'
        f'<a:prstDash val="solid"/></a:{side}>'
    )


def _table(slide, s: Table) -> None:
    n_rows, n_cols = len(s.rows) + 1, len(s.header)
    frame = slide.shapes.add_table(
        n_rows,
        n_cols,
        Inches(MARGIN_X),
        Inches(BODY_TOP),
        Inches(BODY_W),
        Inches(TABLE_ROW_H * n_rows),
    )
    frame.name = "表"
    table = frame.table
    table.first_row = False
    table.horz_banding = False
    table._tbl.tblPr.find(f"{{{_NS_A}}}tableStyleId").text = _TABLE_STYLE_NONE
    col_w = int(Inches(BODY_W) / n_cols)
    for col in table.columns:
        col.width = col_w
    for r in table.rows:
        r.height = Inches(TABLE_ROW_H)
    grid = [s.header, *s.rows]
    for ri, row in enumerate(grid):
        for ci, text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(CELL_MARGIN_V)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.background()
            _fill_paragraph(cell.text_frame.paragraphs[0], text, PT_CELL, GRAY if ri == 0 else INK)
            tc_pr = cell._tc.get_or_add_tcPr()
            underline = ri in (0, n_rows - 1)
            sides = [("lnL", False), ("lnR", False), ("lnT", False), ("lnB", underline)]
            for pos, (side, visible) in enumerate(sides):
                tc_pr.insert(pos, parse_xml(_line_xml(side, visible)))


def _block(slide, b: Block, x: float) -> None:
    if isinstance(b, BulletsBlock):
        tf = _textbox(slide, "箇条書き", x, BODY_TOP, COLUMN_W, BODY_BOX_H)
        _paragraphs(tf, b.items, PT_COLUMN_BULLET, INK, True, gap=10)
    elif isinstance(b, ImageBlock):
        _picture(slide, b.grid, b.alt, x, BODY_TOP)
    elif isinstance(b, TextBlock):
        tf = _textbox(slide, "文", x, BODY_TOP, COLUMN_W, BODY_BOX_H)
        _paragraphs(tf, [b.text], PT_COLUMN_TEXT, INK, False)


def _two_columns(slide, s: TwoColumn) -> None:
    _block(slide, s.left, MARGIN_X)
    _block(slide, s.right, MARGIN_X + COLUMN_W + COLUMN_GAP)


# --- 文書全体 ---


def _set_theme_fonts(prs) -> None:
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")

    def fix(m: re.Match) -> str:
        block = m.group(0)
        block = re.sub(r'<a:latin typeface="[^"]*"', f'<a:latin typeface="{FONT}"', block)
        block = re.sub(r'<a:ea typeface="[^"]*"', f'<a:ea typeface="{FONT}"', block)
        return re.sub(r'(<a:font script="Jpan" typeface=")[^"]*"', rf'\g<1>{FONT}"', block)

    xml = re.sub(r"<a:(major|minor)Font>.*?</a:\1Font>", fix, xml, flags=re.DOTALL)
    # 書式スキームのグラデーション・影・立体は、単色と空の効果に置き換える（個数は規格どおり 3）
    solid = '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>' * 3
    effects = "<a:effectStyle><a:effectLst/></a:effectStyle>" * 3
    for tag, body in (
        ("fillStyleLst", solid),
        ("bgFillStyleLst", solid),
        ("effectStyleLst", effects),
    ):
        xml = re.sub(
            rf"<a:{tag}>.*?</a:{tag}>",
            lambda _, t=tag, b=body: f"<a:{t}>{b}</a:{t}>",
            xml,
            flags=re.DOTALL,
        )
    theme._blob = xml.encode("utf-8")


def _clean_package(prs) -> None:
    """既定テンプレートの雛形（サムネイル、プリンタ設定、4:3 の表記）を残さない。"""
    package = prs.part.package
    for rid, rel in list(package._rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            package._rels.pop(rid)
    for rid, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith("/printerSettings"):
            prs.part.rels.pop(rid)
    for part in package.iter_parts():
        if str(part.partname) == "/docProps/app.xml":
            xml = part.blob.decode("utf-8")
            xml = xml.replace("On-screen Show (4:3)", "Widescreen")
            part._blob = xml.encode("utf-8")


def _set_properties(prs, title: str) -> None:
    cp = prs.core_properties
    now = datetime.now(UTC).replace(tzinfo=None, microsecond=0)
    cp.title = title
    for field in ("author", "last_modified_by", "subject", "keywords", "comments", "category"):
        setattr(cp, field, "")
    cp.created = now
    cp.modified = now


def render(deck: Deck, out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    # テンプレートの 4:3 の種別が残ると PowerPoint の表示が食い違う
    prs._element.sldSz.attrib.pop("type", None)
    _set_theme_fonts(prs)
    for number, s in enumerate(deck.slides, start=1):
        blank = isinstance(s, Statement)
        slide = prs.slides.add_slide(
            prs.slide_layouts[LAYOUT_BLANK if blank else LAYOUT_TITLE_ONLY]
        )
        _white_background(slide)
        if isinstance(s, Cover):
            _cover(slide, s)
        elif isinstance(s, Statement):
            _statement(slide, s)
        else:
            _title(slide, s.heading, MARGIN_X, HEAD_TOP, BODY_W, 0.9, PT_HEADING)
            if isinstance(s, Bullets):
                _bullets(slide, s)
            elif isinstance(s, Figure):
                _figure(slide, s)
            elif isinstance(s, Table):
                _table(slide, s)
            elif isinstance(s, TwoColumn):
                _two_columns(slide, s)
        if number > 1:
            _page_number(slide, number)
    _set_properties(prs, deck.title)
    _clean_package(prs)
    prs.save(str(out_path))
